"""Логика корпоративных предложений: меню, подбор состава, расчёт сумм."""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl

SIZES = ["23 см", "30 см", "35 см", "40 см"]  # 15 см отдельно не продаётся
THIN_SIZES = {"30 см", "35 см", "40 см"}
DOUGHS = ["Традиционное", "Тонкое"]

# Острые пиццы — рядом с ними в КП ставится перчик. Дополняйте при необходимости.
SPICY = {"Большая Бонанза", "Мексиканская"}

# Шаблоны тарифов: пицца -> «вес» (доля в заказе). На 60 пиццах дают ровно пример КП.
TIER_TEMPLATES = {
    "Базовый": {
        "Пепперони": 15, "Ветчина и грибы": 12, "Супер Папа": 9,
        "Мясная": 9, "Маргарита": 9, "Цыпленок Барбекю": 6,
    },
    "Оптимальный": {
        "Пепперони": 12, "Супер Папа": 9, "Ветчина и грибы": 9, "Цыпленок Барбекю": 6,
        "Большая Бонанза": 6, "Цыпленок Рэнч": 6, "Мясная": 6, "Маргарита": 6,
    },
    "Премиальный": {
        "Пепперони": 12, "Супер Папа": 12, "Большая Бонанза": 6,
        "Итальянская с моцареллой и пепперони": 6, "Баварская": 6, "Мясное барбекю": 6,
        "Цыпленок Рэнч": 6, "8 сыров": 3, "Мясная": 3,
    },
}

DEFAULT_FOOTER = (
    "Если есть другие предпочтения по составу или хочется что-то поменять — дайте знать. "
    "Если есть сумма, в которую нужно уложиться, — соберём предложение исходя из бюджета и пожеланий."
)


# ---------------------------------------------------------------- меню
@dataclass
class Menu:
    pizzas: dict  # name -> {"prices": {size: int|None}, "thin": bool}
    extras: list  # [{"cat","name","size","price","label"}]

    def pizza_price(self, name: str, size: str):
        p = self.pizzas.get(name)
        return None if p is None else p["prices"].get(size)

    def eligible(self, size: str, dough: str) -> list[str]:
        """Пиццы, которые есть в этом размере и тесте."""
        out = []
        for name, p in self.pizzas.items():
            if p["prices"].get(size) is None:
                continue
            if dough == "Тонкое" and (not p["thin"] or size not in THIN_SIZES):
                continue
            out.append(name)
        return out

    def extra_by_label(self, label: str):
        for e in self.extras:
            if e["label"] == label:
                return e
        return None


def _num(v):
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str):
        digits = "".join(ch for ch in v if ch.isdigit())
        return int(digits) if digits else None
    return None


def load_menu(path_or_file) -> Menu:
    wb = openpyxl.load_workbook(path_or_file, data_only=True)
    ws = wb["Пиццы"]
    header = [c.value for c in ws[1]]
    pizzas = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]
        if not name or not isinstance(name, str) or name.startswith(("Примечания", "•")):
            continue
        if all(v in (None, "") for v in row[1:6]):
            continue
        prices = {}
        for i, h in enumerate(header):
            if isinstance(h, str) and h.endswith("см") and h in SIZES:
                v = row[i]
                prices[h] = v if isinstance(v, (int, float)) else None
        thin_col = header.index("Тонкое тесто") if "Тонкое тесто" in header else None
        thin = str(row[thin_col]).strip().lower() == "да" if thin_col is not None else False
        pizzas[name.strip()] = {"prices": prices, "thin": thin}

    extras = []
    if "Остальное меню" in wb.sheetnames:
        ws2 = wb["Остальное меню"]
        cat = ""
        for row in ws2.iter_rows(min_row=2, values_only=True):
            name, size, price = row[0], row[1], row[2]
            if not name:
                continue
            if price in (None, "") and (size in (None, "")):
                cat = str(name).strip()  # строка-раздел: «Закуски», «Напитки»...
                continue
            p = _num(price)
            if p is None:
                continue
            size = (size or "").strip() if isinstance(size, str) else ""
            label = f"{name.strip()}" + (f", {size}" if size else "")
            extras.append({"cat": cat, "name": name.strip(), "size": size, "price": p,
                           "label": label, "desc": (row[4] or "") if len(row) > 4 else ""})
    return Menu(pizzas=pizzas, extras=extras)


# ---------------------------------------------------------------- распределение
def distribute(total: int, weights: list[float], step: int = 1) -> list[int]:
    """Раскладывает total по весам методом наибольших остатков, кратно step."""
    if not weights or total <= 0:
        return [0] * len(weights)
    step = max(1, int(step))
    units, rest = divmod(total, step)
    s = float(sum(weights)) or 1.0
    raw = [units * w / s for w in weights]
    base = [math.floor(r) for r in raw]
    left = units - sum(base)
    order = sorted(range(len(raw)), key=lambda i: (raw[i] - base[i], weights[i]), reverse=True)
    for i in order[:left]:
        base[i] += 1
    out = [b * step for b in base]
    if rest:  # то, что не делится на кратность, — к самой большой позиции
        out[max(range(len(out)), key=lambda i: out[i])] += rest
    return out


def fill_template(tier: str, menu: Menu, size: str, dough: str, total: int, step: int = 1):
    """Состав по шаблону тарифа. Недоступные в этом размере/тесте пиццы пропускаются."""
    ok = set(menu.eligible(size, dough))
    tpl = [(n, w) for n, w in TIER_TEMPLATES.get(tier, {}).items() if n in ok]
    if not tpl:
        return []
    qty = distribute(total, [w for _, w in tpl], step)
    return [(n, q) for (n, _), q in zip(tpl, qty) if q > 0]


def fill_even(names: list[str], total: int, step: int = 1):
    qty = distribute(total, [1] * len(names), step)
    return [(n, q) for n, q in zip(names, qty) if q > 0]


def optimize(menu: Menu, size: str, dough: str, total: int, *,
             kinds_min: int = 1, kinds_max: int = 6, step: int = 1,
             budget: float | None = None, discount: float = 0.0,
             extras_cost: float = 0.0, extras_discounted: bool = False,
             goal: str = "min", allowed: list[str] | None = None,
             max_share: float | None = None):
    """Целочисленный подбор состава.

    goal="min"  — самый дешёвый набор;
    goal="max"  — максимально близко к бюджету (не превышая его).
    budget — сумма к оплате со скидкой (вместе с закусками/напитками).
    Возвращает (список (пицца, кол-во), сообщение об ошибке | None).
    """
    from scipy.optimize import Bounds, LinearConstraint, milp

    names = [n for n in (allowed or menu.eligible(size, dough))
             if menu.pizza_price(n, size) is not None]
    if not names:
        return [], "Нет подходящих пицц для этого размера и теста."
    step = max(1, int(step))
    if total % step:
        return [], f"Общее количество ({total}) не делится на кратность {step}."
    units = total // step
    kinds_max = min(kinds_max, len(names), units)
    kinds_min = min(max(1, kinds_min), kinds_max)
    # равномернее: одна позиция не больше чем ~total/kinds_min (с запасом)
    if max_share is None:
        cap_units = math.ceil(units / kinds_min)
    else:
        cap_units = max(1, math.ceil(units * max_share))

    n = len(names)
    price = np.array([menu.pizza_price(nm, size) for nm in names], dtype=float)
    k = 1 - discount / 100.0
    # переменные: z_i (единиц по step) и y_i (вид используется)
    c_price = np.concatenate([price * step, np.zeros(n)])
    popular = set().union(*[set(t) for t in TIER_TEMPLATES.values()])
    pop_bonus = np.array([30.0 * step if nm in popular else 0.0 for nm in names])
    if goal == "max":
        # ближе к бюджету; при прочих равных — популярные пиццы
        c = -c_price - np.concatenate([pop_bonus, np.zeros(n)])
    else:
        # при равной цене чуть предпочитаем больше видов
        c = c_price - np.concatenate([np.zeros(n), np.full(n, 0.01)])
    A, lo, hi = [], [], []
    # сумма пицц
    A.append(np.concatenate([np.full(n, 1.0), np.zeros(n)])); lo.append(units); hi.append(units)
    # z_i <= cap*y_i ; z_i >= y_i
    for i in range(n):
        r = np.zeros(2 * n); r[i] = 1; r[n + i] = -cap_units; A.append(r); lo.append(-np.inf); hi.append(0)
        r = np.zeros(2 * n); r[i] = 1; r[n + i] = -1; A.append(r); lo.append(0); hi.append(np.inf)
    # число видов
    A.append(np.concatenate([np.zeros(n), np.ones(n)])); lo.append(kinds_min); hi.append(kinds_max)
    # бюджет
    if budget:
        extra = extras_cost * (k if extras_discounted else 1)
        A.append(np.concatenate([price * step * k, np.zeros(n)]))
        lo.append(-np.inf); hi.append(budget - extra)
    integrality = np.ones(2 * n)
    bounds = Bounds(np.zeros(2 * n), np.concatenate([np.full(n, cap_units), np.ones(n)]))
    res = milp(c, constraints=LinearConstraint(np.array(A), lo, hi),
               integrality=integrality, bounds=bounds, options={"time_limit": 10})
    if res.x is None:
        if budget:
            return [], "Не удалось уложиться в бюджет с такими условиями — увеличьте бюджет или уменьшите число видов."
        return [], "Не удалось подобрать состав с такими условиями."
    z = np.round(res.x[:n]).astype(int)
    items = [(names[i], int(z[i] * step)) for i in range(n) if z[i] > 0]
    items.sort(key=lambda t: (-t[1], menu.pizza_price(t[0], size)))
    return items, None


# ---------------------------------------------------------------- расчёт
def tier_totals(tier: dict) -> dict:
    """tier = {"pizzas":[{name,qty,price}], "extras":[{label,qty,price}], "discount":..,
    "extras_discounted": bool}"""
    pz = sum(int(r["qty"]) * int(r["price"]) for r in tier["pizzas"])
    ex = sum(int(r["qty"]) * int(r["price"]) for r in tier.get("extras", []))
    cnt = sum(int(r["qty"]) for r in tier["pizzas"])
    d = float(tier.get("discount", 0)) / 100
    disc_base = pz + (ex if tier.get("extras_discounted") else 0)
    to_pay = round(disc_base * (1 - d)) + (0 if tier.get("extras_discounted") else ex)
    return {"pizza_sum": pz, "extras_sum": ex, "count": cnt,
            "full": pz + ex, "to_pay": to_pay}


def rub(v) -> str:
    return f"{int(round(v)):,}".replace(",", " ") + " ₽"
