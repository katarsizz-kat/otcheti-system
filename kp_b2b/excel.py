"""Excel-версия КП: редактируемый файл с формулами + скрытый лист с данными,
чтобы через несколько дней загрузить файл обратно на сайт и продолжить правку."""
from __future__ import annotations

import io
import json
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .core import Menu

GREEN, DARK, LIME, CREAM, ALT = "0D3D26", "083625", "C9E265", "F4E8DC", "FBF6EF"
DATA_SHEET = "_data"
EMPTY_ROWS = 3  # пустые строки в каждом блоке, чтобы дописать позиции
FMT_RUB = '#,##0 "₽"'

_f = lambda **k: Font(name="Arial", **k)
_fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
_thin = Side(style="thin", color="E6DDD2")
_thick = Side(style="medium", color=GREEN)


def build_xlsx(kp: dict, menu: Menu) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "КП"
    mn = wb.create_sheet("Меню")
    dt = wb.create_sheet(DATA_SHEET)

    # --- лист «Меню»: цены на момент формирования КП
    size = kp["size"]
    mn["A1"], mn["B1"], mn["D1"], mn["E1"] = f"Пицца ({size})", "Цена", "Закуски / напитки", "Цена"
    pz_rows = sorted(menu.eligible(size, kp["dough"]))
    saved_pz = {r["name"]: r["price"] for t in kp["tiers"] for r in t["pizzas"] if r.get("name")}
    names = sorted(set(pz_rows) | set(saved_pz))
    for i, n in enumerate(names, start=2):
        mn.cell(i, 1, n)
        mn.cell(i, 2, saved_pz.get(n, menu.pizza_price(n, size)))
    saved_ex = {r["label"]: r["price"] for t in kp["tiers"] for r in t.get("extras", []) if r.get("label")}
    labels = [e["label"] for e in menu.extras]
    labels += [l for l in saved_ex if l not in labels]
    for i, l in enumerate(labels, start=2):
        e = menu.extra_by_label(l)
        mn.cell(i, 4, l)
        mn.cell(i, 5, saved_ex.get(l, e["price"] if e else None))
    for c in "ABDE":
        mn[f"{c}1"].font = _f(bold=True)
    mn.column_dimensions["A"].width = 40
    mn.column_dimensions["D"].width = 48
    pz_last, ex_last = len(names) + 1, len(labels) + 1
    dv_pz = DataValidation(type="list", formula1=f"=Меню!$A$2:$A${pz_last}", allow_blank=True)
    dv_ex = DataValidation(type="list", formula1=f"=Меню!$D$2:$D${ex_last}", allow_blank=True)
    dv_pz.showErrorMessage = dv_ex.showErrorMessage = False
    ws.add_data_validation(dv_pz)
    ws.add_data_validation(dv_ex)

    # --- лист «КП»
    ws.sheet_view.showGridLines = False
    widths = {"A": 3, "B": 44, "C": 10, "D": 13, "E": 15}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.column_dimensions["A"].hidden = True  # служебные метки для обратной загрузки

    def band(r, text, color=GREEN, font=None, height=None):
        for c in range(2, 6):
            ws.cell(r, c).fill = _fill(color)
        ws.cell(r, 2, text).font = font or _f(bold=True, color="FFFFFF", size=11)
        if height:
            ws.row_dimensions[r].height = height

    r = 1
    band(r, kp.get("title") or "Предложение для корпоративного заказа",
         font=_f(bold=True, color="FFFFFF", size=15), height=30)
    r += 1
    info = [("client", "Клиент", kp.get("client", "")),
            ("size", "Размер", size), ("dough", "Тесто", kp["dough"]),
            ("addr", "Адреса доставки",
             "; ".join(f"{a['addr']} — {a['qty']}" if a.get("qty") else a["addr"]
                       for a in kp.get("addresses", []) if a.get("addr")))]
    for key, lbl, val in info:
        ws.cell(r, 1, key)
        ws.cell(r, 2, lbl).font = _f(color="6B7A70")
        ws.cell(r, 3, val).font = _f(bold=True)
        r += 1
    r += 1

    def header(r, first):
        for c, t in enumerate([first, "Кол-во", "Цена", "Сумма"], start=2):
            cell = ws.cell(r, c, t)
            cell.font = _f(bold=True, color="FFFFFF", size=9)
            cell.fill = _fill(GREEN)
            cell.alignment = Alignment(horizontal="left" if c == 2 else "right")

    def item_row(r, name, qty, rng, zebra):
        ws.cell(r, 2, name)
        ws.cell(r, 3, qty if qty else None)
        ws.cell(r, 4, f'=IF(B{r}="","",IFERROR(VLOOKUP(B{r},{rng},2,FALSE),""))')
        ws.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),"",C{r}*D{r})')
        for c in range(2, 6):
            cell = ws.cell(r, c)
            cell.font = _f(size=10)
            cell.border = Border(bottom=_thin)
            if zebra:
                cell.fill = _fill(ALT)
        ws.cell(r, 3).alignment = Alignment(horizontal="right")
        ws.cell(r, 4).number_format = ws.cell(r, 5).number_format = FMT_RUB

    def total_row(r, label, qty_formula, sum_formula, key):
        ws.cell(r, 1, key)
        ws.cell(r, 2, label)
        if qty_formula:
            ws.cell(r, 3, qty_formula)
        ws.cell(r, 5, sum_formula)
        for c in range(2, 6):
            cell = ws.cell(r, c)
            cell.font = _f(bold=True, color=GREEN)
            cell.border = Border(top=_thick)
        ws.cell(r, 5).number_format = FMT_RUB

    for t in kp["tiers"]:
        ws.cell(r, 1, "tier")
        band(r, t["name"], height=22)
        ws.cell(r, 4, "Скидка, %").font = _f(color=LIME, bold=True)
        ws.cell(r, 4).alignment = Alignment(horizontal="right")
        dcell = ws.cell(r, 5, float(t.get("discount", 0)))
        dcell.font = _f(bold=True, color=GREEN)
        dcell.fill = _fill(LIME)
        dcell.alignment = Alignment(horizontal="center")
        d_ref = f"$E${r}"
        r += 1
        ws.cell(r, 1, "pz_head")
        header(r, f"Пицца ({size}{', тонкое' if kp['dough'] == 'Тонкое' else ''})")
        r += 1
        start = r
        rows = [x for x in t["pizzas"] if x.get("name")] + [{}] * EMPTY_ROWS
        for i, x in enumerate(rows):
            item_row(r, x.get("name"), x.get("qty"), f"Меню!$A:$B", i % 2)
            dv_pz.add(f"B{r}")
            r += 1
        end = r - 1
        pz_sum = f"SUM(E{start}:E{end})"
        total_row(r, "Итого пицца (без скидки)", f"=SUM(C{start}:C{end})", f"={pz_sum}", "pz_total")
        pz_total_ref = f"E{r}"
        r += 2
        ws.cell(r, 1, "ex_head")
        header(r, "Закуски и напитки")
        r += 1
        start = r
        rows = [x for x in t.get("extras", []) if x.get("label")] + [{}] * EMPTY_ROWS
        for i, x in enumerate(rows):
            item_row(r, x.get("label"), x.get("qty"), f"Меню!$D:$E", i % 2)
            dv_ex.add(f"B{r}")
            r += 1
        end = r - 1
        total_row(r, "Итого закуски и напитки", None, f"=SUM(E{start}:E{end})", "ex_total")
        ex_total_ref = f"E{r}"
        r += 1
        ws.cell(r, 1, "exd")
        ws.cell(r, 2, "Скидка на закуски и напитки (1 — да, 0 — нет)").font = _f(size=9, color="6B7A70")
        ws.cell(r, 5, 1 if t.get("extras_discounted") else 0).alignment = Alignment(horizontal="center")
        exd = f"E{r}"
        r += 1
        ws.cell(r, 2, "Всего без скидки").font = _f(bold=True)
        ws.cell(r, 5, f"={pz_total_ref}+{ex_total_ref}").number_format = FMT_RUB
        ws.cell(r, 5).font = _f(bold=True)
        r += 1
        band(r, "К оплате", color=GREEN, font=_f(bold=True, color=LIME, size=12), height=22)
        pay = ws.cell(r, 5, f"=ROUND(({pz_total_ref}+IF({exd}=1,{ex_total_ref},0))*(1-{d_ref}/100),0)"
                             f"+IF({exd}=1,0,{ex_total_ref})")
        pay.number_format = FMT_RUB
        pay.font = _f(bold=True, color=LIME, size=12)
        r += 3

    ws.cell(r, 1, "footer")
    ws.cell(r, 2, kp.get("footer", "")).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 45
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # --- скрытый лист с полным состоянием КП
    payload = json.dumps(dict(kp, saved_at=datetime.now().isoformat(timespec="minutes")),
                         ensure_ascii=False)
    for i in range(0, len(payload), 30000):
        dt.cell(i // 30000 + 1, 1, payload[i:i + 30000])
    dt.sheet_state = "hidden"
    mn.sheet_state = "visible"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_xlsx(file, menu: Menu) -> dict:
    """Восстанавливает КП из Excel. Правки, сделанные прямо в листе «КП»
    (названия, количество, скидка, клиент), имеют приоритет над сохранёнными данными."""
    wb = openpyxl.load_workbook(file)  # формулы нам не нужны — читаем вводимые значения
    if DATA_SHEET not in wb.sheetnames:
        raise ValueError("Это не файл КП, созданный на сайте (нет служебного листа).")
    dt = wb[DATA_SHEET]
    payload = "".join(str(c.value) for c in dt["A"] if c.value)
    kp = json.loads(payload)
    ws = wb["КП"]

    saved_pz = {r["name"]: r["price"] for t in kp["tiers"] for r in t["pizzas"] if r.get("name")}
    saved_ex = {r["label"]: (r["price"], r.get("cat", "")) for t in kp["tiers"]
                for r in t.get("extras", []) if r.get("label")}
    menu_sheet = wb["Меню"] if "Меню" in wb.sheetnames else None
    sheet_pz, sheet_ex = {}, {}
    if menu_sheet:
        for a, b, _, d, e in menu_sheet.iter_rows(min_row=2, max_col=5, values_only=True):
            if a and isinstance(b, (int, float)):
                sheet_pz[a] = int(b)
            if d and isinstance(e, (int, float)):
                sheet_ex[d] = int(e)

    tiers, cur, section = [], None, None
    for row in ws.iter_rows(min_row=1, max_col=5):
        key = row[0].value
        name, qty = row[1].value, row[2].value
        if key == "client":
            kp["client"] = row[2].value or ""
        elif key == "tier":
            cur = {"name": str(name or "").strip() or f"Вариант {len(tiers) + 1}",
                   "discount": float(row[4].value or 0), "pizzas": [], "extras": [],
                   "extras_discounted": False}
            tiers.append(cur)
            section = None
        elif key == "pz_head":
            section = "pz"
        elif key == "ex_head":
            section = "ex"
        elif key in ("pz_total", "ex_total"):
            section = None
        elif key == "exd" and cur is not None:
            cur["extras_discounted"] = str(row[4].value).strip() in ("1", "1.0", "да", "True")
        elif key == "footer":
            kp["footer"] = name or ""
        elif section and cur is not None and name and qty not in (None, ""):
            try:
                q = int(float(qty))
            except (TypeError, ValueError):
                continue
            if q <= 0:
                continue
            name = str(name).strip()
            if section == "pz":
                price = sheet_pz.get(name) or saved_pz.get(name) or menu.pizza_price(name, kp["size"]) or 0
                cur["pizzas"].append({"name": name, "qty": q, "price": int(price)})
            else:
                e = menu.extra_by_label(name)
                price = sheet_ex.get(name) or (saved_ex.get(name) or (None,))[0] or (e["price"] if e else 0)
                cat = (saved_ex.get(name) or (None, ""))[1] or (e["cat"] if e else "")
                cur["extras"].append({"label": name, "qty": q, "price": int(price), "cat": cat})
    for t in tiers:  # одинаковые позиции, дописанные отдельной строкой, — складываем
        for part, key in (("pizzas", "name"), ("extras", "label")):
            merged = {}
            for r in t[part]:
                if r[key] in merged:
                    merged[r[key]]["qty"] += r["qty"]
                else:
                    merged[r[key]] = dict(r)
            t[part] = list(merged.values())
    if tiers:
        kp["tiers"] = tiers
    return kp
