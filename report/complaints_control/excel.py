"""
Генерация 4-листового Excel-отчёта «Контроль жалоб» — ТЗ §6.

Листы:
- "Чаты"        : сопоставление чатов (Тюмень, затем СПб) + «не подтверждено источником»;
- "issues_report": сопоставление issues_report (Тюмень, затем СПб) + «не подтверждено источником»;
- "Свод"        : полный список «НЕ ЗАНЕСЕНО», % незанесённых по сотруднику, по ресторанам;
- "Метрики"     : 5 метрик с текстовым описанием перед каждой таблицей.

Цветовая индикация — заливка строки ПО статусу + отдельный текстовый
столбец статуса (оба одновременно, как решил заказчик).
"""
from __future__ import annotations

import io
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import constants as c

THIN = Side(style="thin")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FONT = Font(name="Calibri", size=14, bold=True)
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="2C3E50")
DESC_FONT = Font(name="Calibri", size=10, italic=True, color="555555")
HEADER_FONT = Font(name="Calibri", size=10, bold=True)
DATA_FONT = Font(name="Calibri", size=10)

HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
HEADER_FONT_WHITE = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

STATUS_FILLS = {
    status: PatternFill("solid", fgColor=color)
    for status, color in c.STATUS_FILL_COLORS.items()
}


def _clean(val):
    if val is None:
        return ""
    try:
        if pd.isna(val):
            return ""
    except Exception:
        pass
    return val


def _set_widths(ws, columns, sample_rows) -> None:
    for j, colname in enumerate(columns, start=1):
        name = str(colname)
        width = len(name) + 2
        for row in sample_rows:
            try:
                width = max(width, len(str(_clean(row.iloc[j - 1]))) + 2)
            except Exception:
                pass
        if "Текст" in name or "Комментарий" in name:
            width = 60
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 12), 70)


def _write_title(ws, row: int, text: str, font: Font = TITLE_FONT) -> int:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = font
    return row + 1


def _write_status_table(ws, start_row: int, title: str, df: pd.DataFrame) -> int:
    """Таблица сопоставления: заливка строки по колонке «Статус» + сама
    колонка «Статус» остаётся видимым текстом (ТЗ §6)."""
    row = _write_title(ws, start_row, title, SECTION_FONT)
    if df is None or df.empty:
        ws.cell(row=row, column=1, value="Нет данных").font = DATA_FONT
        return row + 2

    headers = list(df.columns)
    status_col_idx = headers.index("Статус") + 1 if "Статус" in headers else None

    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=j, value=str(h))
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    r = row + 1
    for _, data_row in df.iterrows():
        status = str(data_row.get("Статус", "")) if status_col_idx else None
        fill = STATUS_FILLS.get(status)
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=j, value=_clean(data_row[h]))
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT if j == 1 or "Текст" in str(h) else CENTER
            if fill is not None:
                cell.fill = fill
        r += 1

    _set_widths(ws, headers, [x[1] for x in df.head(80).iterrows()])
    return r + 2


def _write_plain_table(ws, start_row: int, title: str, df: pd.DataFrame, description: Optional[str] = None) -> int:
    row = _write_title(ws, start_row, title, SECTION_FONT)
    if description:
        ws.cell(row=row, column=1, value=description).font = DESC_FONT
        row += 1
    if df is None or df.empty:
        ws.cell(row=row, column=1, value="Нет данных").font = DATA_FONT
        return row + 2

    headers = list(df.columns)
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=j, value=str(h))
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    r = row + 1
    for _, data_row in df.iterrows():
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=j, value=_clean(data_row[h]))
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = LEFT if j == 1 else CENTER
        r += 1

    _set_widths(ws, headers, [x[1] for x in df.head(80).iterrows()])
    return r + 2


METRIC_DESCRIPTIONS = {
    "unpaid": (
        "Что показывает: среди жалоб с оформленным решением о компенсации — "
        "какая доля реально не выплачена (по сотруднику и по ресторану). "
        "Зачем нужна: находит, где гости не получают обещанную компенсацию."
    ),
    "deviation": (
        "Что показывает: доля подтверждённых нарушений стандарта («Отклонение» = «Да») "
        "среди заполненных строк. Только СПб — поля нет в тюменской таблице. "
        "Заполняется нерегулярно — при малом числе строк выводы делать осторожно."
    ),
    "cost": (
        "Что показывает: сумма и среднее по стоимости компенсаций — по ресторану и "
        "по типу жалобы. Зачем нужна: находит точки/категории с наибольшими издержками."
    ),
    "backlog": (
        "Что показывает: доля обращений issues_report в статусе «В ожидании»/«Открыт»/"
        "«В очереди» по очереди/отделу («Исполнитель»). Зачем нужна: показывает, где "
        "скапливаются необработанные обращения."
    ),
    "decision": (
        "Что показывает: как часто каждый ресторан использует каждый способ компенсации "
        "(Баллы/Купон/Возврат/довоз/Замена/без компенсации). Зачем нужна: проверка на "
        "единообразие политики между точками."
    ),
}


def build_complaints_control_excel(
    comparison_chat_tmn: pd.DataFrame,
    comparison_chat_spb: pd.DataFrame,
    not_confirmed_chat_tmn: pd.DataFrame,
    not_confirmed_chat_spb: pd.DataFrame,
    comparison_issues_tmn: pd.DataFrame,
    comparison_issues_spb: pd.DataFrame,
    not_confirmed_issues_tmn: pd.DataFrame,
    not_confirmed_issues_spb: pd.DataFrame,
    not_entered_all: pd.DataFrame,
    employee_share: pd.DataFrame,
    restaurant_summary: pd.DataFrame,
    metric_unpaid: pd.DataFrame,
    metric_deviation: pd.DataFrame,
    metric_cost: pd.DataFrame,
    metric_backlog: pd.DataFrame,
    metric_decision: pd.DataFrame,
    period_label: str = "",
) -> io.BytesIO:
    wb = Workbook()

    # ------------------------------------------------------
    # ЛИСТ 1: ЧАТЫ
    # ------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Чаты"
    row = 1
    if period_label:
        row = _write_title(ws1, row, period_label)
        row += 1
    row = _write_status_table(ws1, row, "Тюмень — сопоставление (чат)", comparison_chat_tmn)
    row = _write_status_table(ws1, row, "Тюмень — не подтверждено источником", not_confirmed_chat_tmn)
    row = _write_status_table(ws1, row, "Санкт-Петербург — сопоставление (чат)", comparison_chat_spb)
    row = _write_status_table(ws1, row, "Санкт-Петербург — не подтверждено источником", not_confirmed_chat_spb)

    # ------------------------------------------------------
    # ЛИСТ 2: issues_report
    # ------------------------------------------------------
    ws2 = wb.create_sheet("issues_report")
    row = 1
    if period_label:
        row = _write_title(ws2, row, period_label)
        row += 1
    row = _write_status_table(ws2, row, "Тюмень — сопоставление (issues_report)", comparison_issues_tmn)
    row = _write_status_table(ws2, row, "Тюмень — не подтверждено источником", not_confirmed_issues_tmn)
    row = _write_status_table(ws2, row, "Санкт-Петербург — сопоставление (issues_report)", comparison_issues_spb)
    row = _write_status_table(ws2, row, "Санкт-Петербург — не подтверждено источником", not_confirmed_issues_spb)

    # ------------------------------------------------------
    # ЛИСТ 3: СВОД
    # ------------------------------------------------------
    ws3 = wb.create_sheet("Свод")
    row = 1
    if period_label:
        row = _write_title(ws3, row, period_label)
        row += 1
    row = _write_status_table(ws3, row, "НЕ ЗАНЕСЕНО — все источники", not_entered_all)
    row = _write_plain_table(ws3, row, "% незанесённых по сотруднику (только чаты)", employee_share)
    row = _write_plain_table(ws3, row, "По ресторанам", restaurant_summary)

    # ------------------------------------------------------
    # ЛИСТ 4: МЕТРИКИ
    # ------------------------------------------------------
    ws4 = wb.create_sheet("Метрики")
    row = 1
    if period_label:
        row = _write_title(ws4, row, period_label)
        row += 1
    row = _write_plain_table(ws4, row, "1. Доля невыплаченной компенсации", metric_unpaid, METRIC_DESCRIPTIONS["unpaid"])
    row = _write_plain_table(ws4, row, "2. Доля подтверждённых нарушений стандарта", metric_deviation, METRIC_DESCRIPTIONS["deviation"])
    row = _write_plain_table(ws4, row, "3. Стоимость компенсаций", metric_cost, METRIC_DESCRIPTIONS["cost"])
    row = _write_plain_table(ws4, row, "4. Незакрытый бэклог issues_report", metric_backlog, METRIC_DESCRIPTIONS["backlog"])
    row = _write_plain_table(ws4, row, "5. Распределение способов компенсации", metric_decision, METRIC_DESCRIPTIONS["decision"])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
